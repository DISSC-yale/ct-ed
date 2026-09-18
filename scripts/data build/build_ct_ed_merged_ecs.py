"""
build_ct_merged_with_ecs.py
============================
Extends the base merged CT education panel with ECS (Education Cost Sharing)
grant shell data for FY2024–FY2027.

Inputs
------
- ct_education_merged.csv          (output of build_ct_merged.py, same directory)
- FY_2023-2024_ECS_shell.xlsx      )
- FY_2024-2025_ECS_shell.xlsx      )  ECS shell spreadsheets
- FY_2025-2026_ECS_shell.xlsx      )  (place in same directory as this script)
- FY_2026-2027_ECS_shell_and_out_years.xlsx  )

Output
------
- ct_ed_merged_ecs.csv

Dependencies: pandas, openpyxl
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# config

HERE = os.path.dirname(os.path.abspath(__file__))

MERGED_CSV = os.path.join(HERE, "ct_education_merged.csv")
OUT_CSV    = os.path.join(HERE, "ct_ed_merged_ecs.csv")

SHELLS = {
    2024: (os.path.join(HERE, "FY_2023-2024_ECS_shell.xlsx"),                  "FY 24"),
    2025: (os.path.join(HERE, "FY_2024-2025_ECS_shell.xlsx"),                  "FY 25"),
    2026: (os.path.join(HERE, "FY_2025-2026_ECS_shell.xlsx"),                  "FY 26"),
    2027: (os.path.join(HERE, "FY_2026-2027_ECS_shell_and_out_years.xlsx"),    "FY 27"),
}

# ECS shell column map
# The shells are messy multi-row-header Excel files. District data begins at
# row 27; the town name is always in column index 8. All four fiscal years use
# the same column layout (only header label text changes year to year, not the
# column positions). Column indices not listed here are blank spacer columns.
#
# We read every data column — including intermediate formula-output columns —
# and give each a stable snake_case name that strips the year-specific text
# from the original headers (e.g. "(10/2022)" vs "(10/2023)" become the same
# column name across all four years).

COL_MAP = {
    0:  "drg",
    1:  "psd_flag",
    2:  "alliance_district_flag",
    3:  "alliance_non_reform_flag",
    4:  "reform_district_flag",
    5:  "wealth_decile",
    6:  "pic_decile",
    7:  "town_code",
    # col 8 = town name; used as join key, not kept as a data column
    10: "resident_students",
    12: "frpl_students",
    13: "frpl_poverty_portion",                    # col 2 × 0.30
    14: "net_resident_students_60pct",             # col 1 × 0.60
    15: "excess_resident_students",                # max(col2 - col4, 0)
    16: "excess_resident_students_portion",        # col 5 × 0.15
    18: "frpl_pct_above_60pct",
    19: "concentrated_poverty_students",
    21: "ell_students",
    22: "ell_need_students_portion",               # col 7 × 0.25
    23: "frpl_weight",                             # col 2 × item A
    24: "total_need_students",                     # col 1+3+6+8
    25: "engl_3yr_avg",                            # 3-yr avg equalized net grand list
    26: "total_population",
    27: "engl_per_capita",                         # col 10 / col 11
    28: "engl_adjustment_factor",                  # EEPC / threshold
    29: "median_household_income",
    30: "mhi_adjustment_factor",                   # MHI / threshold
    31: "wealth_adjustment_factor",
    32: "base_aid_ratio",
    33: "pic_add_base_aid_ratio_adjustment",
    34: "final_base_aid_ratio",
    35: "students_sent_to_regional_district",
    36: "n_regional_district_grades",
    37: "regional_district_bonus",
    38: "students_sent_to_endowed_academies",
    39: "n_endowed_academy_grades",
    40: "endowed_academy_bonus",
    41: "base_formula_aid",
    42: "fully_funded_grant",
    43: "ecs_actual_fy2017",                       # FY2017 ECS actual; baseline reference
    44: "fully_funded_grant_with_alliance_hh",
    45: "prior_year_entitlement",
    46: "grant_adjustment",
    47: "fully_funded_gt_prior_year",              # Yes/No flag
    48: "phase_in_amount",
    49: "ecs_entitlement_without_alliance_hh",
    50: "ecs_entitlement_with_alliance_hh",        # ← primary policy output
    51: "change_from_prior_year",                  # present in FY2024 only; NaN in FY2025-27
}

DATA_COL_MAP = {k: v for k, v in COL_MAP.items() if k != 8}   # exclude town name col

# All ECS columns in the output are prefixed "ecs__"
PREFIX = "ecs__"


# helpers

def normalize_name(s):
    """
    Strip common district-name suffixes for join-key matching.
    e.g. "Hartford School District" -> "HARTFORD"
         "Andover"                  -> "ANDOVER"
    """
    s = str(s).strip().upper()
    for suffix in (
        " SCHOOL DISTRICT", " SCHOOL DIST", " SCHOOL",
        " DISTRICT", " DIST",
        " PUBLIC SCHOOLS", " CITY SCHOOLS",
    ):
        if s.endswith(suffix):
            s = s[: -len(suffix)].strip()
    return s


def read_shell(path, sheet_name, fiscal_year):
    """
    Read one ECS shell workbook/sheet and return a tidy DataFrame.

    - Opens with data_only=True so computed cells return their cached values,
      not the formula strings.
    - District rows run from row 27 to row 195 (169 towns).
    - Town name is always at column index 8.
    - Blank/formula-string values are stored as NaN.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]

    rows = []
    for row in ws.iter_rows(min_row=27, max_row=195, values_only=True):
        town_name = row[8] if len(row) > 8 else None
        if not isinstance(town_name, str) or not town_name.strip():
            continue  # skip blank/header bleed-through rows

        rec = {"town_name": town_name.strip(), "fiscal_year": fiscal_year}
        for col_idx, col_name in DATA_COL_MAP.items():
            val = row[col_idx] if len(row) > col_idx else None
            # data_only=True should give computed values, but guard defensively
            if isinstance(val, str) and val.startswith("="):
                val = None
            rec[col_name] = val
        rows.append(rec)

    wb.close()
    return pd.DataFrame(rows)


# main

# 1. Load the base merged panel
print(f"Loading {MERGED_CSV} ...")
panel = pd.read_csv(MERGED_CSV)
panel["district_code"] = panel["district_code"].astype("Int64")
print(f"  {len(panel):,} rows × {len(panel.columns)} columns")

# 2. Build name --> district_code and name --> district_name lookups from the panel.
#    These are used to resolve ECS shell town names (bare, e.g. "Andover") to
#    the EdSight LEA codes and canonical district names in the merged panel.
name_to_code = (
    panel[["district_name", "district_code"]]
    .dropna()
    .drop_duplicates("district_name")
    .assign(norm=lambda df: df["district_name"].map(normalize_name))
    .set_index("norm")["district_code"]
    .to_dict()
)
name_to_distname = (
    panel[["district_name"]]
    .dropna()
    .drop_duplicates("district_name")
    .assign(norm=lambda df: df["district_name"].map(normalize_name))
    .set_index("norm")["district_name"]
    .to_dict()
)

# 3. Read all four ECS shells and stack into one long DataFrame
print("\nReading ECS shells ...")
shell_frames = []
for fy, (path, sheet) in SHELLS.items():
    df = read_shell(path, sheet, fy)
    print(f"  FY{fy}: {len(df)} rows")
    shell_frames.append(df)

shells = pd.concat(shell_frames, ignore_index=True)

# 4. Resolve district_code for each shell row via normalized name lookup.
#    20 of the 169 shell towns are member towns of Regional School Districts
#    that have no standalone LEA code in EdSight. These will not match and
#    will have district_code = NaN. They are kept as new rows carrying only
#    ECS data (no enrollment, spending, assessment, etc. counterpart exists
#    in the merged panel for them). See the documentation for details.
shells["_norm"] = shells["town_name"].map(normalize_name)
shells["district_code"] = shells["_norm"].map(name_to_code).astype("Int64")
shells = shells.drop(columns=["_norm"])

n_matched   = shells["district_code"].notna().sum()
n_unmatched = shells["district_code"].isna().sum()
unmatched_towns = (
    shells[shells["district_code"].isna()]["town_name"].unique().tolist()
)
print(f"\n  Matched to district_code: {n_matched}")
print(f"  Unmatched (RSD member towns, new rows): {n_unmatched}")
print(f"  Unmatched towns: {sorted(unmatched_towns)}")

# 5. Prefix all ECS data columns with "ecs__"
data_cols = list(DATA_COL_MAP.values())
shells = shells.rename(
    columns={c: f"{PREFIX}{c}" for c in data_cols}
)
shells = shells.rename(columns={"town_name": f"{PREFIX}town_name"})
ecs_cols = [c for c in shells.columns if c.startswith(PREFIX)]

# 6. Merge ECS columns into existing panel rows (left join on district_code × fiscal_year).
#    This adds ECS data to the FY2024/2025/2026 rows already present in the panel
#    for the 149 matched town districts.
print("\nMerging ECS data into existing panel rows ...")
shells_for_merge = shells[["district_code", "fiscal_year"] + ecs_cols].copy()
panel = panel.merge(shells_for_merge, on=["district_code", "fiscal_year"], how="left")

# 7. Append new rows for the 20 unmatched RSD-member towns (all four years).
#    These rows carry ECS data only; all other columns will be NaN.
unmatched_rows = shells[shells["district_code"].isna()].copy()
unmatched_rows["district_code"] = pd.NA
# Set district_name from the bare town name so the row is identifiable
unmatched_rows.insert(0, "district_name", unmatched_rows[f"{PREFIX}town_name"])

# 8. Append new rows for the FY2027 matched towns.
#    FY2027 has no enrollment data (and no rows in the base panel), so these
#    are entirely new rows. Resolve the canonical district_name from the lookup.
fy27_matched = shells[
    (shells["fiscal_year"] == 2027) & shells["district_code"].notna()
].copy()
fy27_matched["district_name"] = (
    fy27_matched[f"{PREFIX}town_name"]
    .map(normalize_name)
    .map(name_to_distname)
)

panel = pd.concat([panel, unmatched_rows, fy27_matched], ignore_index=True)

# 9. Sort and assert no duplicates on the (district_code, fiscal_year) key
#    for rows that have a district_code (unmatched towns have NaN, which is fine).
panel = panel.sort_values(
    ["district_code", "fiscal_year"], na_position="last"
).reset_index(drop=True)

dups = (
    panel.dropna(subset=["district_code"])
    .duplicated(["district_code", "fiscal_year"])
    .sum()
)
assert dups == 0, f"ERROR: {dups} duplicate (district_code, fiscal_year) pairs found."

# 10. Write output
print(f"\nWriting {OUT_CSV}")
panel.to_csv(OUT_CSV, index=False)

print(f"\nDone.")
print(f"  Rows    : {len(panel):,}")
print(f"  Columns : {len(panel.columns):,}  (+{len(ecs_cols)} ECS columns)")
print(f"  Districts (with code): {panel['district_code'].nunique():,}")
print(f"  New ECS-only rows (unmatched towns): {len(unmatched_rows):,}")
print(f"  New FY2027 rows: {len(fy27_matched):,}")
print(f"  Year range: {panel['fiscal_year'].min()}–{panel['fiscal_year'].max()}")
